"""Strict, bounded XLSX ingestion. Unknown columns (including email) are discarded."""
from io import BytesIO
from zipfile import ZipFile, BadZipFile
import re
from openpyxl import load_workbook

FIELDS = {
    "id": ["attendee id", "attendee code", "id"],
    "name": ["display name", "name", "full name"],
    "offers": ["skills i can help with", "offers", "skills"],
    "needs": ["what i want help with", "needs", "help needed"],
    "interests": ["interests", "topics of interest"],
    "project": ["what i am building", "project", "what i'm building"],
    "window": ["available networking window", "availability", "window"],
    "consent": ["allow my profile to appear in event matches", "consent"],
}
REQUIRED = {"id", "name", "offers", "needs", "window", "consent"}

def normalize(value):
    return re.sub(r"\s+", " ", str(value or "").strip().lower().replace("_", " "))

def read_xlsx(data: bytes):
    if len(data) > 5 * 1024 * 1024:
        raise ValueError("File exceeds 5 MB. Upload a smaller XLSX export.")
    try:
        with ZipFile(BytesIO(data)) as archive:
            if sum(x.file_size for x in archive.infolist()) > 25 * 1024 * 1024:
                raise ValueError("Workbook expands beyond the 25 MB limit.")
        book = load_workbook(BytesIO(data), read_only=True, data_only=False)
    except (BadZipFile, KeyError, OSError) as exc:
        raise ValueError("This is not a readable XLSX workbook.") from exc
    try:
        sheet = book["Form Responses 1"] if "Form Responses 1" in book.sheetnames else book.worksheets[0]
        if (sheet.max_column or 0) > 50 or (sheet.max_row or 0) > 1001:
            raise ValueError("Use at most 50 columns and 1,000 attendees, with headers in row 1.")
        rows = sheet.iter_rows()
        headers = next(rows, ())
        if len(headers) > 50:
            raise ValueError("Use at most 50 columns.")
        mapping = {}
        for i, cell in enumerate(headers):
            for field, aliases in FIELDS.items():
                if normalize(cell.value) in aliases:
                    if field in mapping:
                        raise ValueError(f"Duplicate column for {field}.")
                    mapping[field] = i
        if missing := REQUIRED - mapping.keys():
            raise ValueError("Missing columns: " + ", ".join(sorted(missing)))
        profiles, seen, skipped = [], set(), 0
        for row_num, row in enumerate(rows, 2):
            if row_num > 1001 or len(row) > 50:
                raise ValueError("Use at most 50 columns and 1,000 attendees.")
            if all(c.value is None for c in row):
                continue
            values = {}
            for field, i in mapping.items():
                cell = row[i] if i < len(row) else None
                if cell is not None and cell.data_type == "f":
                    raise ValueError(f"Row {row_num}: formulas are not accepted in profile fields.")
                value = str((cell.value if cell is not None else "") or "").strip()
                if len(value) > 1000:
                    raise ValueError(f"Row {row_num}: {field} exceeds 1,000 characters.")
                values[field] = value
            # Duplicate identities are ambiguous even when one row declines consent.
            # Never keep an older consenting row over a later opt-out response.
            row_id = values["id"].upper()
            if row_id and row_id in seen:
                raise ValueError(f"Row {row_num}: duplicate attendee ID {row_id}.")
            if row_id:
                seen.add(row_id)
            # Consent is explicit; blanks and unrecognized values never imply consent.
            if normalize(values["consent"]) not in {"yes", "true", "i agree"}:
                skipped += 1
                continue
            for field in REQUIRED - {"consent", "needs"}:
                if not values[field]:
                    raise ValueError(f"Row {row_num}: {field} is empty.")
            if not re.fullmatch(r"[A-Za-z0-9_-]{1,32}", values["id"]):
                raise ValueError(f"Row {row_num}: attendee ID must be 1–32 letters, numbers, underscores or hyphens.")
            values["id"] = values["id"].upper()
            profiles.append({k: values.get(k, "") for k in FIELDS if k != "consent"})
        return profiles, skipped
    finally:
        book.close()
