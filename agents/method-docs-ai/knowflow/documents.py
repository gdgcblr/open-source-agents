from pathlib import Path
from pypdf import PdfReader
from docx import Document

SUPPORTED = {".txt", ".md", ".pdf", ".docx"}

def read_document(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix in {".txt", ".md"}:
        return path.read_text(encoding="utf-8", errors="replace")
    if suffix == ".pdf":
        reader = PdfReader(str(path))
        return "\n".join(page.extract_text() or "" for page in reader.pages)
    if suffix == ".docx":
        doc = Document(str(path))
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                paragraphs.append(" | ".join(cell.text.strip() for cell in row.cells))
        return "\n".join(paragraphs)
    raise ValueError(f"Unsupported document type: {suffix}")

def discover_documents(root: Path) -> list[Path]:
    if not root.exists():
        return []
    return sorted(
        p for p in root.rglob("*")
        if p.is_file() and p.suffix.lower() in SUPPORTED
        and not any(part.startswith(".") for part in p.relative_to(root).parts)
    )

def chunk_text(text: str, chunk_size: int = 1200, overlap: int = 180) -> list[str]:
    text = "\n".join(line.strip() for line in text.replace("\x00", "").splitlines() if line.strip())
    if not text:
        return []
    if chunk_size <= 0 or overlap < 0 or overlap >= chunk_size:
        raise ValueError("Require chunk_size > overlap >= 0.")
    chunks = []
    start = 0
    while start < len(text):
        end = min(start + chunk_size, len(text))
        if end < len(text):
            # Prefer ending at a sentence/paragraph boundary when it is nearby.
            boundary = max(text.rfind("\n", start, end), text.rfind(". ", start, end))
            if boundary > start + int(chunk_size * 0.55):
                end = boundary + 1
        piece = text[start:end].strip()
        if piece:
            chunks.append(piece)
        if end >= len(text):
            break
        start = max(end - overlap, start + 1)
    return chunks
