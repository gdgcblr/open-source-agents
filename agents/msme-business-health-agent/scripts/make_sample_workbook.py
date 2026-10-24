"""Generate a SYNTHETIC, deliberately messy 4-tab workbook. No real data.

Run: python scripts/make_sample_workbook.py
"""
import os
import random
from datetime import datetime, timedelta

from openpyxl import Workbook

random.seed(42)
OUT = os.path.join("data", "sample_workbook.xlsx")
os.makedirs("data", exist_ok=True)
wb = Workbook()

# ---- Sales: title rows, mixed date formats, blanks, duplicates, totals row
ws = wb.active
ws.title = "Sales"
ws["A1"] = "Demo Traders - Sales Register"
ws.append([])
ws.append(["Date", "Invoice No", "Customer", "SKU", "Qty", "Amount"])
rows = []
start = datetime(2026, 9, 1)
for i in range(60):
    d = start + timedelta(days=random.randint(0, 40))
    date_val = d.strftime("%d/%m/%Y") if i % 3 else d  # mix strings and real dates
    amt = None if i % 11 == 0 else random.randint(500, 9000)
    rows.append([date_val, f"INV{1000+i}", f"Customer {random.randint(1,12)}", f"SKU{random.randint(1,8)}",
                 random.randint(1, 20), amt])
rows[5] = ["31/02/2026", "INV1999", "Customer 3", "SKU2", 4, 1200]  # invalid date
rows += [rows[0], rows[1], rows[2]]  # duplicates
for r in rows:
    ws.append(r)
ws.append(["Total", None, None, None, None, 123456])

# ---- Inventory
ws = wb.create_sheet("Inventory")
ws.append(["SKU", "Item", "Stock", "Reorder Level", "Last Restocked"])
for i in range(1, 9):
    ws.append([f"SKU{i}", f"Item {i}", random.randint(0, 120), 20, f"{random.randint(1,28):02d}/09/2026"])

# ---- Support tickets
ws = wb.create_sheet("Support Tickets")
ws.append(["Ticket ID", "Customer", "Issue", "Opened", "Resolved"])
for i in range(25):
    o = start + timedelta(days=random.randint(0, 35))
    res = None if i % 6 == 0 else (o + timedelta(days=random.randint(1, 9))).strftime("%d/%m/%Y")
    ws.append([f"T{i+1}", f"Customer {random.randint(1,12)}", random.choice(["Late delivery", "Damaged", "Billing"]),
               o.strftime("%d/%m/%Y"), res])

# ---- Order events (event log for the bottleneck X-ray later)
ws = wb.create_sheet("Order Events")
ws.append(["order_id", "activity", "timestamp", "user"])
acts = ["Order Received", "Payment Confirmed", "Packed", "Dispatched", "Delivered"]
for i in range(1, 31):
    t = datetime(2026, 9, 1, 10, 0) + timedelta(hours=random.randint(0, 600))
    for a in acts:
        ws.append([f"ORD{i:03d}", a, t.strftime("%d/%m/%Y %H:%M"), random.choice(["Ravi", "Meena", "Auto"])])
        gap = random.randint(1, 6)
        if a == "Packed":
            gap = random.randint(20, 70)  # planted bottleneck: Packed -> Dispatched
        t += timedelta(hours=gap)

wb.save(OUT)
print("Wrote", OUT)
